import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


class FillBlankSection:

    def __init__(
            self,
            title="Fill in the Blanks",
            questions=None,
            formatting=None
    ):

        self.title = title

        # List of original questions
        # Example:
        # "He (kicked) the ball."
        self.questions = questions if questions else []

        # For future font toolbar support
        self.formatting = formatting or {
            "size": 18,
            "bold": False,
            "italic": False,
            "underline": False
        }

    # ----------------------------------------------------
    # QUESTION HELPERS
    # ----------------------------------------------------

    def add_question(self, text):

        text = text.strip()

        if text:
            self.questions.append(text)

    def get_preview(self, question):
        """
        Used by the frontend to display previews.
        """

        answers = re.findall(r'\((.*?)\)', question)

        worksheet_text = re.sub(
            r'\((.*?)\)',
            "_____",
            question
        )

        if answers:
            answer_text = ", ".join(answers)
            return f"{worksheet_text} ({answer_text})"

        return worksheet_text

    def get_size(self):

        return len(self.questions) * 2

    # ----------------------------------------------------
    # DOCX HELPERS
    # ----------------------------------------------------

    def set_cell_border(self, cell, **kwargs):

        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        tcBorders = OxmlElement('w:tcBorders')

        for edge in ('top', 'bottom', 'left', 'right'):

            if edge in kwargs:

                edge_data = kwargs.get(edge)

                tag = OxmlElement(f'w:{edge}')

                tag.set(
                    qn('w:val'),
                    edge_data.get('val', 'single')
                )

                tag.set(
                    qn('w:sz'),
                    str(edge_data.get('sz', 6))
                )

                tag.set(qn('w:space'), '0')
                tag.set(qn('w:color'), 'auto')

                tcBorders.append(tag)

        tcPr.append(tcBorders)

    def set_row_height_exact(self, row, height):

        tr = row._tr
        trPr = tr.get_or_add_trPr()

        trHeight = OxmlElement('w:trHeight')

        trHeight.set(qn('w:val'), str(height))
        trHeight.set(qn('w:hRule'), 'exact')

        trPr.append(trHeight)

    def add_writing_lines(self, cell):

        p = cell.paragraphs[0]
        p._element.getparent().remove(p._element)

        inner_table = cell.add_table(rows=3, cols=1)

        for i, row in enumerate(inner_table.rows):

            inner_cell = row.cells[0]

            inner_cell.text = ""

            p = inner_cell.paragraphs[0]

            run = p.add_run(" ")
            run.font.size = Pt(8)

            if i == 0:

                self.set_cell_border(
                    inner_cell,
                    top={"val": "single", "sz": 8},
                    bottom={"val": "dashed", "sz": 8}
                )

            elif i == 1:

                self.set_cell_border(
                    inner_cell,
                    bottom={"val": "single", "sz": 8}
                )

            elif i == 2:

                self.set_cell_border(
                    inner_cell,
                    bottom={"val": "single", "sz": 8}
                )

            self.set_row_height_exact(row, 200)

    # ----------------------------------------------------
    # EXPORTS
    # ----------------------------------------------------

    def export(self, doc):

        if not self.questions:
            return

        for i, sentence in enumerate(
                self.questions,
                start=1):

            worksheet = re.sub(
                r'\((.*?)\)',
                "__________",
                sentence
            )

            table = doc.add_table(rows=2, cols=1)

            # Question row
            sentence_cell = table.rows[0].cells[0]

            sentence_cell.text = (
                f"{i}. {worksheet}"
            )

            p = sentence_cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)

            # Writing row
            writing_cell = table.rows[1].cells[0]

            self.add_writing_lines(writing_cell)

            doc.add_paragraph("")

    def export_answers(self, doc):

        for i, sentence in enumerate(
                self.questions,
                start=1):

            answers = re.findall(
                r'\((.*?)\)',
                sentence
            )

            answer_text = ", ".join(answers)

            doc.add_paragraph(
                f"{i}. {answer_text}"
            )

    def get_title(self):

        return self.title