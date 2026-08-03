import random

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


class UnscrambleSection:

    def __init__(
            self,
            title="Unscramble the Words",
            words=None,
            show_word_bank=True,
            formatting=None
    ):

        self.title = title

        self.words = words if words else []

        self.show_word_bank = show_word_bank

        # For future toolbar formatting support
        self.formatting = formatting or {
            "size": 18,
            "bold": False,
            "italic": False,
            "underline": False
        }

    def get_words(self):
        """
        Remove empty values and convert to lowercase.
        """

        return [
            word.strip().lower()
            for word in self.words
            if word.strip()
        ]

    def scramble_word(self, word):

        if len(word) <= 1:
            return word

        scrambled = word

        while scrambled == word:
            letters = list(word)
            random.shuffle(letters)
            scrambled = "".join(letters)

        return scrambled

    def get_scrambled_words(self):

        words = self.get_words()

        return [
            self.scramble_word(word)
            for word in words
        ]

    def get_size(self):

        words = self.get_scrambled_words()

        rows = (len(words) + 1) // 2

        return rows * 2

    # ----------------------------------------------------
    # DOCX HELPERS
    # ----------------------------------------------------

    def set_cell_border(self, cell, **kwargs):

        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        tcBorders = OxmlElement('w:tcBorders')

        for edge in ('top', 'bottom', 'left', 'right'):

            if edge in kwargs:

                edge_data = kwargs.get(edge)

                tag = OxmlElement(f'w:{edge}')

                tag.set(
                    qn('w:val'),
                    edge_data.get('val', 'single')
                )

                tag.set(
                    qn('w:sz'),
                    str(edge_data.get('sz', 6))
                )

                tag.set(qn('w:space'), '0')
                tag.set(qn('w:color'), 'auto')

                tcBorders.append(tag)

        tcPr.append(tcBorders)

    def set_row_height_exact(self, row, height):

        tr = row._tr
        trPr = tr.get_or_add_trPr()

        trHeight = OxmlElement('w:trHeight')

        trHeight.set(qn('w:val'), str(height))
        trHeight.set(qn('w:hRule'), 'exact')

        trPr.append(trHeight)

    def add_writing_lines(self, cell):

        p = cell.paragraphs[0]
        p._element.getparent().remove(p._element)

        inner_table = cell.add_table(rows=3, cols=1)

        for i, row in enumerate(inner_table.rows):

            inner_cell = row.cells[0]

            inner_cell.text = ""

            p = inner_cell.paragraphs[0]

            run = p.add_run(" ")
            run.font.size = Pt(8)

            if i == 0:

                self.set_cell_border(
                    inner_cell,
                    top={"val": "single", "sz": 8},
                    bottom={"val": "dashed", "sz": 8}
                )

            elif i == 1:

                self.set_cell_border(
                    inner_cell,
                    bottom={"val": "single", "sz": 8}
                )

            elif i == 2:

                self.set_cell_border(
                    inner_cell,
                    bottom={"val": "single", "sz": 8}
                )

            self.set_row_height_exact(row, 200)

    # ----------------------------------------------------
    # EXPORTS
    # ----------------------------------------------------

    def export(self, doc):

        original_words = self.get_words()
        words = self.get_scrambled_words()

        if not words:
            return

        # Word Bank
        if self.show_word_bank:

            sorted_words = sorted(original_words)

            word_bank_table = doc.add_table(rows=2, cols=6)

            word_index = 0

            for r in range(2):

                for c in range(6):

                    if word_index < len(sorted_words):

                        cell = word_bank_table.rows[r].cells[c]

                        paragraph = cell.paragraphs[0]

                        paragraph.alignment = (
                            WD_PARAGRAPH_ALIGNMENT.CENTER
                        )

                        cell.text = sorted_words[word_index]

                        border_args = {}

                        if r == 0:
                            border_args["top"] = {
                                "val": "single",
                                "sz": 8
                            }

                        if r == 1:
                            border_args["bottom"] = {
                                "val": "single",
                                "sz": 8
                            }

                        if c == 0:
                            border_args["left"] = {
                                "val": "single",
                                "sz": 8
                            }

                        if c == 5:
                            border_args["right"] = {
                                "val": "single",
                                "sz": 8
                            }

                        self.set_cell_border(
                            cell,
                            **border_args
                        )

                        word_index += 1

            doc.add_paragraph()

        rows = (len(words) + 1) // 2

        table = doc.add_table(
            rows=rows,
            cols=4
        )

        for i, word in enumerate(words):

            row = i % rows
            col = (i // rows) * 2

            table.rows[row].cells[col].text = (
                f"{i + 1}. {word}"
            )

            answer_cell = table.rows[row].cells[col + 1]

            self.add_writing_lines(answer_cell)

    def export_answers(self, doc):

        words = self.get_words()

        for i, word in enumerate(words, start=1):

            doc.add_paragraph(f"{i}. {word}")

    def get_title(self):

        return self.title