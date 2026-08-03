import random
import string
import copy

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


class WordSearchSection:

    def __init__(
            self,
            title,
            difficulty,
            grid_size,
            instructions,
            words,
            formatting):

        self.title = title
        self.difficulty = difficulty
        self.grid_size = grid_size
        self.instructions = instructions
        self.formatting = formatting
        self.words = [
            w.strip().upper()
            for w in words
            if w.strip()
        ]

    def get_title(self):
        return self.title

    def get_grid_size(self):

        size_map = {
            "Small (8 x 8)": 8,
            "Medium (12 x 12)": 12,
            "Large (15 x 15)": 15,
            "Extra Large (20 x 20)": 20
        }

        return size_map.get(
            self.grid_size,
            8
        )

    def get_directions(self):

        if self.difficulty == "Beginner":

            return [
                (1, 0),   # down
                (0, 1)    # right
            ]

        elif self.difficulty == "Medium":

            return [
                (1, 0),
                (0, 1),
                (1, 1),
                (-1, 1)
            ]

        else:

            return [
                (1, 0),
                (0, 1),
                (-1, 0),
                (0, -1),
                (1, 1),
                (-1, 1),
                (1, -1),
                (-1, -1)
            ]

    def create_grid(self):

        size = self.get_grid_size()

        return [
            ["" for _ in range(size)]
            for _ in range(size)
        ]

    def place_word(self, grid, word):

        size = len(grid)

        for _ in range(100):

            dr, dc = random.choice(
                self.get_directions()
            )

            row = random.randint(
                0,
                size - 1
            )

            col = random.randint(
                0,
                size - 1
            )

            end_row = row + dr * (len(word) - 1)
            end_col = col + dc * (len(word) - 1)

            if not (
                0 <= end_row < size and
                0 <= end_col < size
            ):
                continue

            fits = True

            for i in range(len(word)):

                r = row + dr * i
                c = col + dc * i

                if grid[r][c] not in (
                        "",
                        word[i]
                ):
                    fits = False
                    break

            if fits:

                for i in range(len(word)):

                    r = row + dr * i
                    c = col + dc * i

                    grid[r][c] = word[i]

                return True

        return False

    def fill_empty(self, grid):

        for r in range(len(grid)):
            for c in range(len(grid)):

                if grid[r][c] == "":

                    grid[r][c] = random.choice(
                        string.ascii_uppercase
                    )

    def generate_puzzle(self):

        grid = self.create_grid()

        size = self.get_grid_size()

        valid_words = [
            w for w in self.words
            if len(w) <= size
        ]

        for word in valid_words:

            self.place_word(
                grid,
                word
            )

        self.solution_grid = copy.deepcopy(grid)

        self.fill_empty(grid)

        return grid

    def export(self, doc):

        if self.instructions:

            p = doc.add_paragraph()
            run = p.add_run(
                self.instructions
            )
            run.italic = True

        grid = self.generate_puzzle()

        table = doc.add_table(
            rows=len(grid),
            cols=len(grid)
        )

        for r in range(len(grid)):
            for c in range(len(grid)):

                cell = table.rows[r].cells[c]

                cell.text = grid[r][c]

                cell.paragraphs[0].alignment = (
                    WD_PARAGRAPH_ALIGNMENT.CENTER
                )

        if self.words:

            doc.add_paragraph()

            p = doc.add_paragraph()

            run = p.add_run("Words\n")
            run.bold = True

            p.add_run(
                "   ".join(
                    sorted(self.words)
                )
            )

    def export_answers(self, doc):

        doc.add_paragraph(
            "Word Search Answer Key"
        )

        grid = getattr(
            self,
            "solution_grid",
            None
        )

        if not grid:

            self.generate_puzzle()
            grid = self.solution_grid

        table = doc.add_table(
            rows=len(grid),
            cols=len(grid)
        )

        for r in range(len(grid)):
            for c in range(len(grid)):

                letter = grid[r][c]

                table.rows[r].cells[c].text = (
                    letter if letter else "."
                )

    def get_title(self):
        return self.title