from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class TranslationSection:

    def __init__(self, title, pairs, direction, formatting):

        self.title = title
        self.pairs = pairs or []
        self.direction = direction
        self.formatting = formatting


        self.translation_cache = {}

    def get_translations(self):
        translations = []

        for pair in self.pairs:
            source = (pair.get("source") or "").strip()
            target = (pair.get("target") or "").strip()

            if source or target:
                translations.append((source, target))

        return translations

    def set_cell_border(self, cell, **kwargs):

        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        tcBorders = OxmlElement('w:tcBorders')

        for edge in ('top', 'bottom', 'left', 'right'):

            if edge in kwargs:

                edge_data = kwargs.get(edge)

                tag = OxmlElement(f'w:{edge}')

                tag.set(
                    qn('w:val'),
                    edge_data.get('val', 'single')
                )

                tag.set(
                    qn('w:sz'),
                    str(edge_data.get('sz', 6))
                )

                tag.set(qn('w:space'), '0')
                tag.set(qn('w:color'), 'auto')

                tcBorders.append(tag)

        tcPr.append(tcBorders)

    def set_row_height_exact(self, row, height):

        tr = row._tr
        trPr = tr.get_or_add_trPr()

        trHeight = OxmlElement('w:trHeight')

        trHeight.set(qn('w:val'), str(height))
        trHeight.set(qn('w:hRule'), 'exact')

        trPr.append(trHeight)

    def add_writing_lines(self, cell):

        p = cell.paragraphs[0]

        p._element.getparent().remove(
            p._element
        )

        inner_table = cell.add_table(
            rows=3,
            cols=1
        )

        for i, row in enumerate(inner_table.rows):

            inner_cell = row.cells[0]

            inner_cell.text = ""

            p = inner_cell.paragraphs[0]

            run = p.add_run(" ")
            run.font.size = Pt(8)

            if i == 0:

                self.set_cell_border(
                    inner_cell,
                    top={"val": "single", "sz": 8},
                    bottom={"val": "dashed", "sz": 8}
                )

            elif i == 1:

                self.set_cell_border(
                    inner_cell,
                    bottom={"val": "single", "sz": 8}
                )

            elif i == 2:

                self.set_cell_border(
                    inner_cell,
                    bottom={"val": "single", "sz": 8}
                )

            self.set_row_height_exact(
                row,
                200
            )

    def export(self, doc):

        translations = self.get_translations()

        if not translations:
            return

        rows = (len(translations) + 1) // 2

        table = doc.add_table(
            rows=rows,
            cols=4
        )

        for i, (english, korean) in enumerate(
            translations
        ):

            if self.direction == "English → Korean":
                question = english
            else:
                question = korean

            row = i % rows
            col = (i // rows) * 2

            table.rows[row].cells[col].text = (
                f"{i + 1}. {question}"
            )

            answer_cell = (
                table.rows[row].cells[col + 1]
            )

            self.add_writing_lines(
                answer_cell
            )

    def export_answers(self, doc):

        translations = self.get_translations()

        for i, (english, korean) in enumerate(
            translations,
            start=1
        ):

            if self.direction == "English → Korean":

                question = english
                answer = korean

            else:

                question = korean
                answer = english

            doc.add_paragraph(
                f"{i}. {question} → {answer}"
            )

    def get_title(self):
        return self.title